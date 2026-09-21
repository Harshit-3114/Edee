"""Bulk student uploads and commercial tracking for the coaching portal."""
import uuid

import pytest_asyncio
from httpx import AsyncClient
from sqlalchemy import text


CSV_OK = (
    "name,email,phone,stream\n"
    "Aarav Sharma,aarav@example.com,9876543210,UG\n"
    "Aarav Twice,aarav@example.com,9876543210,UG\n"
    "Bad Email,not-an-email,9876543211,UG\n"
    "Bad Phone,phone@example.com,123,UG\n"
    "Bad Stream,stream@example.com,9876543212,XX\n"
)


@pytest_asyncio.fixture
async def a_centre(db_session):
    centre_id = uuid.uuid4()
    await db_session.execute(
        text(
            """
            INSERT INTO coaching_centers (id, name, city, state, active, amount_per_lead)
            VALUES (:id, 'Pragati Academy', 'Pune', 'Maharashtra', true, 10000)
            """
        ),
        {"id": centre_id},
    )
    await db_session.commit()
    return str(centre_id)


@pytest_asyncio.fixture
async def as_coach(auth_as, a_centre):
    auth_as("uid-coach", role="coaching", coaching_centre_id=a_centre)
    return a_centre


class TestUploads:
    async def test_template_downloads_as_xlsx(self, client: AsyncClient, as_coach):
        from openpyxl import load_workbook
        import io

        response = await client.get("/coaching/uploads/template")
        assert response.status_code == 200
        assert "spreadsheetml.sheet" in response.headers["content-type"]
        sheet = load_workbook(io.BytesIO(response.content), read_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        assert [str(h).lower() for h in rows[0]] == [
            "name",
            "email",
            "phone",
            "stream",
            "shortlisted",
        ]
        # Example rows show the shape, including the shortlisted column.
        assert rows[1][0] == "Aarav Sharma"
        assert "::" in str(rows[1][4])

    async def test_upload_counts_created_duplicates_and_errors(
        self, client: AsyncClient, as_coach, db_session
    ):
        response = await client.post(
            "/coaching/uploads",
            files={"file": ("leads.csv", CSV_OK.encode(), "text/csv")},
        )
        assert response.status_code == 201
        summary = response.json()
        assert summary["total"] == 5
        assert summary["created"] == 1
        assert summary["duplicates"] == 1
        assert summary["error_count"] == 3

        history = (await client.get("/coaching/uploads")).json()
        assert len(history) == 1
        assert history[0]["records_created"] == 1

        students = (await client.get("/coaching/uploaded-students")).json()
        assert [s["email"] for s in students] == ["aarav@example.com"]

    async def test_reuploading_is_a_duplicate_not_a_second_row(
        self, client: AsyncClient, as_coach
    ):
        first = (
            await client.post(
                "/coaching/uploads",
                files={"file": ("a.csv", CSV_OK.encode(), "text/csv")},
            )
        ).json()
        second = (
            await client.post(
                "/coaching/uploads",
                files={"file": ("a-again.csv", CSV_OK.encode(), "text/csv")},
            )
        ).json()
        assert first["created"] == 1
        assert second["created"] == 0
        assert second["duplicates"] >= 1

    async def test_platform_email_is_already_taken(
        self, client: AsyncClient, auth_as, as_coach, db_session
    ):
        auth_as("uid-signed", role="student")
        await client.post(
            "/students/",
            json={
                "name": "Taken Person",
                "email": "taken@example.com",
                "phone": "9876500010",
                "stream": "UG",
            },
        )
        auth_as("uid-coach", role="coaching", coaching_centre_id=as_coach)
        summary = (
            await client.post(
                "/coaching/uploads",
                files={
                    "file": (
                        "b.csv",
                        "name,email,phone,stream\nTaken Person,taken@example.com,9876500010,UG\n".encode(),
                        "text/csv",
                    )
                },
            )
        ).json()
        assert summary["created"] == 0
        assert summary["duplicates"] == 1

    async def test_signup_with_invite_marks_the_lead(
        self, client: AsyncClient, auth_as, as_coach, db_session, a_centre
    ):
        await client.post(
            "/coaching/uploads",
            files={
                "file": (
                    "c.csv",
                    "name,email,phone,stream\nSoon Signed,soon@example.com,9876543219,UG\n".encode(),
                    "text/csv",
                )
            },
        )
        await db_session.execute(
            text(
                """
                INSERT INTO coaching_invites (id, coaching_center_id, code, max_uses, uses)
                VALUES (:id, :cid, 'LINKME1234', 5, 0)
                """
            ),
            {"id": uuid.uuid4(), "cid": a_centre},
        )
        await db_session.commit()

        auth_as("uid-soon", role="student")
        assert (
            await client.post(
                "/students/",
                json={
                    "name": "Soon Signed",
                    "email": "soon@example.com",
                    "phone": "9876543219",
                    "stream": "UG",
                    "invite_code": "LINKME1234",
                },
            )
        ).status_code == 201

        status = (
            await db_session.execute(
                text(
                    "SELECT status FROM coaching_students WHERE email = 'soon@example.com'"
                )
            )
        ).fetchone()
        assert status[0] == "signed_up"

    async def test_dashboard_shows_outstanding(
        self, client: AsyncClient, as_coach, db_session, a_centre
    ):
        await client.post(
            "/coaching/uploads",
            files={
                "file": (
                    "d.csv",
                    "name,email,phone,stream\nLead One,one@example.com,9876543220,UG\nLead Two,two@example.com,9876543221,PG\n".encode(),
                    "text/csv",
                )
            },
        )
        data = (await client.get("/coaching/dashboard")).json()
        # 2 leads x 10000 paise, no credit yet.
        assert data["total_leads"] == 2
        assert data["outstanding_amount"] == 20000

        await db_session.execute(
            text("UPDATE coaching_centers SET credit_paise = 2500 WHERE id = :cid"),
            {"cid": a_centre},
        )
        await db_session.commit()
        data = (await client.get("/coaching/dashboard")).json()
        assert data["outstanding_amount"] == 17500
